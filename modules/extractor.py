from pathlib import Path

import fitz
from docx import Document
from pptx import Presentation


def extract_pdf_text(file_path: str) -> str:
    text_chunks = []
    with fitz.open(file_path) as pdf_document:
        for page in pdf_document:
            page_text = page.get_text("text")
            if page_text:
                text_chunks.append(page_text)
    return "\n".join(text_chunks).strip()


def extract_docx_text(file_path: str) -> str:
    document = Document(file_path)
    paragraphs = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text.strip())

    for table in document.tables:
        for row in table.rows:
            cell_values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cell_values:
                paragraphs.append(" | ".join(cell_values))

    return "\n".join(paragraphs).strip()


def extract_pptx_text(file_path: str) -> str:
    presentation = Presentation(file_path)
    slide_text = []

    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())

    return "\n".join(slide_text).strip()


def extract_txt_text(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="replace") as file:
        return file.read().strip()


def extract_text_from_file(file_path: str) -> str:
    extension = Path(file_path).suffix.lower()

    if extension == ".pdf":
        return extract_pdf_text(file_path)
    if extension == ".docx":
        return extract_docx_text(file_path)
    if extension == ".pptx":
        return extract_pptx_text(file_path)
    if extension == ".txt":
        return extract_txt_text(file_path)

    raise ValueError(f"Unsupported file type: {extension}")
